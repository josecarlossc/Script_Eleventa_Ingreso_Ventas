#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
eleventa_sim_venta_gui.py (v4.0.2)

- GUI para ejecutar la venta desde un .xlsx
- Selector de archivo (abre en Descargas)
- Opción "Vender todo a Mayoreo (F11 una vez por código)"
- Ingreso de SKU: repite el CÓDIGO por pieza (qty veces)
- Producto Común: Ctrl+P -> desc -> cant -> precio
- Finaliza con F12 y, si se pide, escribe notas con F4

Correcciones v4.0.2:
- Se agregó la función win32_write_notes (faltaba), con espera del diálogo "NOTAS".
- Conexión Win32 por subcadena + fallback por handle (como v4.0.1).
"""

import sys, time, re, unicodedata, threading, queue
from pathlib import Path
from collections import defaultdict

# ---------- Dependencias externas ----------
try:
    import openpyxl
except Exception:
    print("ERROR: Falta 'openpyxl'. Instala:  pip install openpyxl", file=sys.stderr); raise

try:
    import pyautogui
except Exception:
    pyautogui = None

try:
    import pygetwindow as gw
except Exception:
    gw = None

try:
    from pywinauto.application import Application
    from pywinauto.keyboard import send_keys
    from pywinauto import Desktop
except Exception:
    Application = None
    send_keys = None
    Desktop = None

# ---------- GUI (tkinter) ----------
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

# ---------- Utilidades base ----------
def _norm(s: str) -> str:
    if s is None: return ""
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return s.upper()

def focus_window(window_pattern: str, retries=3, sleep=0.5):
    """Trae al frente la ventana cuyo título contenga window_pattern (subcadena)."""
    if gw is None or not window_pattern: return False
    pat = re.compile(re.escape(window_pattern), re.IGNORECASE)
    for _ in range(retries):
        matches = [w for w in gw.getAllWindows() if w.title and pat.search(w.title)]
        if matches:
            win = sorted(matches, key=lambda w: (w.width*w.height), reverse=True)[0]
            try:
                if win.isMinimized: win.restore()
                win.activate(); time.sleep(sleep); return True
            except Exception: time.sleep(sleep)
    return False

def win32_connect(window_pattern: str):
    """
    Conexión robusta a la ventana por subcadena:
    1) title_re = re.compile('.*<pattern>.*', re.I)
    2) Fallback: Desktop().windows() y conectar por handle del mayor window que coincida
    """
    if Application is None:
        print("pywinauto no disponible; instala: pip install pywinauto", file=sys.stderr)
        return None, None
    pat = re.compile(r".*" + re.escape(window_pattern) + r".*", re.I)

    # Intento 1: conectar por regex (subcadena)
    try:
        app = Application(backend="win32").connect(title_re=pat, timeout=3)
        win = app.top_window()
        return app, win
    except Exception:
        pass

    # Intento 2: buscar con Desktop y conectar por handle del candidato más grande
    try:
        if Desktop is not None:
            candidates = []
            for w in Desktop(backend="win32").windows():
                try:
                    title = w.window_text()
                    if title and pat.search(title):
                        r = w.rectangle()
                        area = max(1, (r.right - r.left) * (r.bottom - r.top))
                        candidates.append((area, w.handle))
                except Exception:
                    continue
            if candidates:
                candidates.sort(reverse=True, key=lambda t: t[0])
                h = candidates[0][1]
                app = Application(backend="win32").connect(handle=h, timeout=3)
                win = app.window(handle=h)
                return app, win
    except Exception:
        pass

    return None, None

def win32_find_edit(win):
    try:
        for c in win.descendants():
            cls = ""; raw = ""
            try: cls = c.friendly_class_name()
            except Exception: pass
            try: raw = c.class_name()
            except Exception: pass
            if cls == "Edit" or raw in ("TEdit","Edit"):
                return c
    except Exception:
        pass
    return None

def _robust_enter_with(target, text: str, interval=0.0):
    try:
        target.type_keys(text + "{ENTER}", with_spaces=True, pause=interval); return True
    except Exception: pass
    try:
        target.type_keys(text, with_spaces=True, pause=interval)
        if send_keys: send_keys("{ENTER}"); return True
    except Exception: pass
    if pyautogui:
        try:
            pyautogui.write(str(text)); pyautogui.press("enter"); return True
        except Exception: pass
    return False

def win32_type_code(win, code: str, interval=0.0):
    if win is None: return False
    try:
        win.set_focus()
        ed = win32_find_edit(win)
        if ed:
            ed.set_focus()
            try: ed.click_input()
            except Exception: pass
            return _robust_enter_with(ed, code, interval)
        else:
            return _robust_enter_with(win, code, interval)
    except Exception as e:
        print(f"[win32] Error tipeando código: {e}", file=sys.stderr)
        return False

def wait_top_contains(app, keyword_norm, timeout=3.0):
    t0 = time.time()
    kw = _norm(keyword_norm)
    while time.time() - t0 < timeout:
        try:
            top = app.top_window()
            if kw in _norm(top.window_text() or ""): return top
        except Exception: pass
        time.sleep(0.05)
    return None

def win32_add_common(app, win, desc: str, qty: int, price: float, interval=0.0, debug=False):
    """Abre Producto Común (^P), escribe desc, cant, precio y espera volver a la venta."""
    try:
        win.set_focus()
        if send_keys: send_keys("^p")
        else: return False

        dlg = wait_top_contains(app, "PRODUCTO COMUN", timeout=2.0)
        if debug and dlg: print(f"[COMMON] Dialog: {dlg.window_text()!r}")
        target = dlg if dlg is not None else win

        desc_clean = (desc or "").strip().strip(" -–—") or "Producto Comun"
        _robust_enter_with(target, desc_clean, interval)
        _robust_enter_with(target, str(int(qty)), interval)
        price_str = f"{float(price):.2f}".replace(",", ".")
        _robust_enter_with(target, price_str, interval)

        wait_top_contains(app, "VENTA DE PRODUCTOS", timeout=2.0) or \
        wait_top_contains(app, "VENTA", timeout=2.0)
        return True
    except Exception as e:
        print(f"[win32] Error en Producto Común: {e}", file=sys.stderr)
        return False

def win32_press_f11():
    """Presiona F11 (mayoreo)."""
    try:
        if send_keys:
            send_keys("{F11}")
            return True
    except Exception:
        pass
    try:
        if pyautogui:
            pyautogui.press("f11")
            return True
    except Exception:
        pass
    return False

def win32_write_notes(app, note_lines, interval=0.0, debug=False):
    """
    Abre la ventana de NOTAS (F4), escribe cada línea (Enter por línea) y confirma (Enter).
    Asume que ya estamos en la ventana COBRAR.
    """
    # Garantizar que estamos en COBRAR
    wait_top_contains(app, "COBRAR", timeout=5.0)

    # Abrir diálogo de notas
    try:
        if send_keys: send_keys("{F4}")
        elif pyautogui: pyautogui.press("f4")
    except Exception:
        pass

    # Esperar el diálogo (el título suele contener 'NOTAS')
    dlg = wait_top_contains(app, "NOTAS", timeout=3.0)
    target = dlg if dlg is not None else app.top_window()

    # Escribir líneas
    for line in note_lines:
        try:
            if send_keys:
                send_keys(line, with_spaces=True, pause=interval)
                send_keys("{ENTER}")
            elif pyautogui:
                pyautogui.write(line)
                pyautogui.press("enter")
        except Exception:
            pass
        time.sleep(0.03)

    # Confirmar (Enter)
    try:
        if send_keys: send_keys("{ENTER}")
        elif pyautogui: pyautogui.press("enter")
    except Exception:
        pass
    if debug:
        print("[NOTES] Notas escritas y confirmadas.")

# ---------- Excel helpers ----------
def infer_columns(ws):
    header = None
    for r in ws.iter_rows(min_row=1, max_row=5, values_only=True):
        if r and any(isinstance(c, str) and c.strip() for c in r):
            header = [(c.strip().lower() if isinstance(c, str) else "") for c in r]
            break
    idx = {"code": None, "qty": None, "desc": None, "price": None, "importe": None}
    start_row = 1
    if header:
        start_row = 2
        code_names   = {"codigo","código","code","sku","barcode","codbarras","código de barras","cod"}
        qty_names    = {"cantidad","qty","cant","unidades"}
        desc_names   = {"descripcion","descripción","producto","nombre","detalle"}
        price_names  = {"precio usado","precio","valor","monto","price"}
        importe_names= {"importe","total","subtotal"}
        for i,name in enumerate(header):
            if name in code_names    and idx["code"]    is None: idx["code"]    = i
            if name in qty_names     and idx["qty"]     is None: idx["qty"]     = i
            if name in desc_names    and idx["desc"]    is None: idx["desc"]    = i
            if name in price_names   and idx["price"]   is None: idx["price"]   = i
            if name in importe_names and idx["importe"] is None: idx["importe"] = i
    if idx["code"] is None: idx["code"] = 0
    return idx, start_row

def parse_common_code(text: str, common_prefix="COMUN"):
    s = str(text).strip()
    m = re.match(rf"^(?:{re.escape(common_prefix)}|{re.escape(common_prefix)+'-'})(.*)$", s, re.IGNORECASE)
    if not m: return None, None
    tail = m.group(1).strip("-_ ")
    parts = re.split(r"[-_]", tail) if tail else []
    price = None; qty = None
    if parts:
        p = parts[0]
        if re.fullmatch(r"\d+[.,]\d{1,2}", p):
            price = float(p.replace(",", "."))
        elif re.fullmatch(r"\d{3,}", p):
            price = float(p) / 100.0
        elif re.fullmatch(r"\d+", p) and len(parts) == 1:
            qty = int(p)
        if len(parts) >= 2 and qty is None and re.fullmatch(r"\d+", parts[1]):
            qty = int(parts[1])
    return price, qty

def load_items_from_xlsx(xlsx_path, common_prefix="COMUN"):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=True); ws = wb.active
    idx, start_row = infer_columns(ws)
    items = []
    for row in ws.iter_rows(min_row=start_row, values_only=True):
        if not row: continue
        code_cell = row[idx["code"]] if idx["code"] is not None and idx["code"] < len(row) else None
        if code_cell is None or str(code_cell).strip() == "": continue
        code = str(code_cell).strip()

        desc = None; price = None; qty = 1; importe = None; qty_from_excel = False

        if idx["desc"] is not None and idx["desc"] < len(row):
            d = row[idx["desc"]]; desc = (str(d).strip() if d is not None else None)

        if idx["price"] is not None and idx["price"] < len(row):
            pr = row[idx["price"]]
            try:
                if isinstance(pr, str): pr = pr.replace("$","").replace(",","").strip()
                price = float(pr)
            except Exception: price = None

        if idx["importe"] is not None and idx["importe"] < len(row):
            im = row[idx["importe"]]
            try:
                if isinstance(im, str): im = im.replace("$","").replace(",","").strip()
                importe = float(im)
            except Exception: importe = None

        if idx["qty"] is not None and idx["qty"] < len(row):
            q = row[idx["qty"]]
            try:
                if q is not None and str(q).strip() != "":
                    qty = int(float(q));  qty = qty if qty>0 else 1
                    qty_from_excel = True
            except Exception: qty = 1

        is_common = bool(re.match(rf"^{re.escape(common_prefix)}(?:$|[-_])", code, re.IGNORECASE))
        if is_common:
            p2, q2 = parse_common_code(code, common_prefix)
            if price is None and p2 is not None: price = p2
            if (not qty_from_excel or qty <= 0) and q2 is not None: qty = q2
            if not desc: desc = common_prefix
            if importe is None and price is not None: importe = price * qty
            items.append({"kind":"common","code":code,"desc":desc,"qty":qty,"price":price,"importe":importe})
        else:
            if importe is None and price is not None: importe = price * qty
            items.append({"kind":"sku","code":code,"desc":desc,"qty":qty,"price":price,"importe":importe})
    return items

def build_notes(items, logger=None, debug=False):
    for it in items:
        pr = it.get("price")
        im = it.get("importe")
        if isinstance(pr, (int, float)) and isinstance(im, (int, float)):
            esperado = float(pr) * int(it["qty"])
            if abs(esperado - float(im)) > 0.01 and logger:
                logger(f"[WARN] Importe Excel no coincide para {it['code']}: qty={it['qty']} price={pr} -> {esperado:.2f} vs Excel={im:.2f}")

    agg = defaultdict(lambda: {"qty": 0, "price": None, "desc": ""})
    for it in items:
        key = (it["kind"], it["code"], float(it.get("price") or 0.0))
        agg[key]["qty"]  += int(it["qty"])
        if agg[key]["price"] is None: agg[key]["price"] = it.get("price")
        if not agg[key]["desc"]:      agg[key]["desc"]  = (it.get("desc") or "")

    def _clean_desc(s):
        return re.sub(r"\s+", " ", (s or "").replace("\t"," ")).strip().strip(" -–—")

    note_lines = ["Codigo Descripcion Cantidad Importe"]
    total_qty = 0
    total_importe = 0.0

    for (kind, code, price), v in agg.items():
        desc = _clean_desc(v["desc"]) or "Producto"
        qty  = int(v["qty"])
        if isinstance(price, (int, float)) and price > 0:
            imp = float(price) * qty
        else:
            imp = 0.0
            for it in items:
                if (it["kind"], it["code"], float(it.get("price") or 0.0)) == (kind, code, price):
                    if isinstance(it.get("importe"), (int, float)):
                        imp += float(it["importe"])
        note_lines.append(f"{code} {desc} {qty} {imp:.2f}")
        total_qty     += qty
        total_importe += imp

    note_lines.append(f"Total Cantidad: {total_qty}")
    note_lines.append(f"Total Importe: {total_importe:.2f}")
    if debug and logger:
        logger(f"[NOTES] Consolidado renglones: {len(agg)}  Total qty={total_qty}  Total=${total_importe:.2f}")
    return note_lines, total_qty, total_importe

# ---------- Lógica principal ----------
def run_sale(config, logger, stop_flag):
    xlsx_path = Path(config["xlsx_path"])
    window    = config["window"]
    start_delay = config["start_delay"]
    between_items = config["between_items"]
    after_code_wait = config["after_code_wait"]
    after_items_wait = config["after_items_wait"]
    write_notes_flag = config["write_notes"]
    finalize = config["finalize"]
    debug = config["debug"]
    common_prefix = config["common_prefix"]
    wholesale_all = config["wholesale_all"]  # aplicar F11 una vez por código

    if not xlsx_path.exists():
        raise FileNotFoundError(f"No se encontró el archivo: {xlsx_path}")

    logger("Cargando Excel...")
    items = load_items_from_xlsx(xlsx_path, common_prefix=common_prefix)
    if not items:
        raise RuntimeError("No se detectaron artículos en el Excel.")

    logger(f"Se cargarán {len(items)} artículos. Inicio en {start_delay:.1f}s...")
    for i in range(int(start_delay),0,-1):
        if stop_flag.is_set(): return
        logger(f"... {i}")
        time.sleep(1)

    logger("Enfocando Eleventa...")
    focus_window(window, retries=4)
    app, win = win32_connect(window)
    if win is None:
        raise RuntimeError("No encontré la ventana por Win32. Ajusta el texto del título (subcadena) o ejecuta con los mismos permisos.")

    wholesale_applied = set()  # códigos a los que YA se les aplicó F11 (para no repetir)

    # Ingreso
    for idx_i, it in enumerate(items, 1):
        if stop_flag.is_set(): return
        if debug:
            logger(f"[{idx_i}] {it['kind'].upper()} -> {it['code']}  qty={it['qty']}  desc={it.get('desc')}  price={it.get('price')}  imp={it.get('importe')}")
        if it["kind"] == "common":
            ok = win32_add_common(app, win, it["desc"], it["qty"], (it["price"] or 0.0), interval=0.0, debug=debug)
            if not ok: logger(f"[WARN] No se pudo ingresar Producto Común para {it['code']}.")
            time.sleep(between_items)
        else:
            reps = max(1, int(it["qty"]))
            for _ in range(reps):
                if stop_flag.is_set(): return
                ok = win32_type_code(win, it["code"], 0.0)
                if not ok:
                    logger(f"[WARN] No se pudo tipear código {it['code']}. Reintentando...")
                    time.sleep(0.15)
                    win32_type_code(win, it["code"], 0.0)
                time.sleep(after_code_wait)
            # Mayoreo por código (una sola vez)
            if wholesale_all and it["code"] not in wholesale_applied:
                time.sleep(0.12)
                if win32_press_f11():
                    wholesale_applied.add(it["code"])
                    if debug: logger(f"[MAYOREO] F11 aplicado a {it['code']}")
                else:
                    logger(f"[WARN] No se pudo enviar F11 para {it['code']}")
            time.sleep(between_items)

    # Notas
    note_lines, _, _ = build_notes(items, logger=logger, debug=debug)

    if finalize:
        if stop_flag.is_set(): return
        logger("Abriendo COBRAR (F12)...")
        try:
            if send_keys: send_keys("{F12}")
            elif pyautogui: pyautogui.press("f12")
        except Exception:
            pass
        time.sleep(max(0.0, after_items_wait))
        wait_top_contains(app, "COBRAR", timeout=5.0)

        if write_notes_flag:
            logger("Escribiendo notas (F4)...")
            win32_write_notes(app, note_lines, interval=0.0, debug=debug)

    logger("Listo. Notas generadas" + (" y escritas." if finalize and write_notes_flag else "."))

# ---------- Interfaz ----------
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Eleventa - Simulador de Venta")
        self.geometry("720x520")
        self.minsize(680, 500)

        # Tema
        style = ttk.Style(self)
        if "vista" in style.theme_names():
            style.theme_use("vista")
        else:
            style.theme_use(style.theme_names()[0])
        style.configure("TButton", padding=6)
        style.configure("Accent.TButton", padding=6)
        style.configure("TCheckbutton", padding=6)

        # Vars
        self.file_var = tk.StringVar()
        self.window_var = tk.StringVar(value="eleventa")
        self.mayoreo_var = tk.BooleanVar(value=False)
        self.finalize_var = tk.BooleanVar(value=True)
        self.notes_var = tk.BooleanVar(value=True)
        self.debug_var = tk.BooleanVar(value=True)
        self.start_delay_var = tk.DoubleVar(value=5.0)
        self.between_items_var = tk.DoubleVar(value=0.35)
        self.after_code_wait_var = tk.DoubleVar(value=0.08)
        self.after_items_wait_var = tk.DoubleVar(value=1.2)
        self.common_prefix_var = tk.StringVar(value="COMUN")

        # Layout
        self._build_ui()

        # Threading
        self.worker = None
        self.stop_flag = threading.Event()
        self.log_q = queue.Queue()
        self.after(100, self._drain_log_queue)

    def _build_ui(self):
        frm = ttk.Frame(self, padding=12)
        frm.pack(fill="both", expand=True)

        # Archivo
        row = 0
        ttk.Label(frm, text="Archivo Excel (.xlsx):").grid(row=row, column=0, sticky="w")
        ent = ttk.Entry(frm, textvariable=self.file_var)
        ent.grid(row=row, column=1, sticky="ew", padx=6)
        ttk.Button(frm, text="Buscar...", command=self._browse).grid(row=row, column=2, sticky="ew")
        frm.columnconfigure(1, weight=1)

        row += 1
        ttk.Label(frm, text="Ventana de Eleventa (subcadena del título):").grid(row=row, column=0, sticky="w", pady=(8,0))
        ttk.Entry(frm, textvariable=self.window_var).grid(row=row, column=1, sticky="ew", padx=6, pady=(8,0))
        ttk.Checkbutton(frm, text="Vender todo a Mayoreo (F11 una vez por código)", variable=self.mayoreo_var)\
            .grid(row=row, column=2, sticky="w", pady=(8,0))

        # Opciones
        row += 1
        opts = ttk.Labelframe(frm, text="Opciones", padding=10)
        opts.grid(row=row, column=0, columnspan=3, sticky="ew", pady=8)

        r = 0
        ttk.Checkbutton(opts, text="Finalizar (F12)", variable=self.finalize_var).grid(row=r, column=0, sticky="w")
        ttk.Checkbutton(opts, text="Escribir notas (F4)", variable=self.notes_var).grid(row=r, column=1, sticky="w")
        ttk.Checkbutton(opts, text="Debug", variable=self.debug_var).grid(row=r, column=2, sticky="w")

        r += 1
        ttk.Label(opts, text="Start delay (s):").grid(row=r, column=0, sticky="e", padx=(0,6), pady=(6,0))
        ttk.Spinbox(opts, from_=0, to=30, increment=0.5, textvariable=self.start_delay_var, width=6)\
            .grid(row=r, column=1, sticky="w", pady=(6,0))

        ttk.Label(opts, text="Entre piezas (s):").grid(row=r, column=2, sticky="e", padx=(12,6), pady=(6,0))
        ttk.Spinbox(opts, from_=0, to=2, increment=0.05, textvariable=self.between_items_var, width=6)\
            .grid(row=r, column=3, sticky="w", pady=(6,0))

        r += 1
        ttk.Label(opts, text="Pausa tras código (s):").grid(row=r, column=0, sticky="e", padx=(0,6), pady=(6,0))
        ttk.Spinbox(opts, from_=0, to=1, increment=0.01, textvariable=self.after_code_wait_var, width=6)\
            .grid(row=r, column=1, sticky="w", pady=(6,0))

        ttk.Label(opts, text="Antes de F12 (s):").grid(row=r, column=2, sticky="e", padx=(12,6), pady=(6,0))
        ttk.Spinbox(opts, from_=0, to=3, increment=0.1, textvariable=self.after_items_wait_var, width=6)\
            .grid(row=r, column=3, sticky="w", pady=(6,0))

        r += 1
        ttk.Label(opts, text="Prefijo 'Producto Común':").grid(row=r, column=0, sticky="e", padx=(0,6), pady=(6,0))
        ttk.Entry(opts, textvariable=self.common_prefix_var, width=10).grid(row=r, column=1, sticky="w", pady=(6,0))

        for c in range(4):
            opts.columnconfigure(c, weight=1)

        # Botones
        row += 1
        btns = ttk.Frame(frm)
        btns.grid(row=row, column=0, columnspan=3, sticky="ew", pady=(4,8))
        self.start_btn = ttk.Button(btns, text="Iniciar venta", style="Accent.TButton", command=self._start)
        self.start_btn.pack(side="left")
        self.stop_btn = ttk.Button(btns, text="Cancelar", command=self._stop, state="disabled")
        self.stop_btn.pack(side="left", padx=6)

        self.pb = ttk.Progressbar(btns, mode="indeterminate")
        self.pb.pack(side="right", fill="x", expand=True)

        # Log
        row += 1
        logf = ttk.Labelframe(frm, text="Consola", padding=6)
        logf.grid(row=row, column=0, columnspan=3, sticky="nsew")
        self.txt = tk.Text(logf, height=12, wrap="word")
        self.txt.pack(fill="both", expand=True)
        frm.rowconfigure(row, weight=1)

    def _log(self, msg: str):
        ts = time.strftime("%H:%M:%S")
        self.log_q.put(f"[{ts}] {msg}\n")

    def _drain_log_queue(self):
        try:
            while True:
                line = self.log_q.get_nowait()
                self.txt.insert("end", line)
                self.txt.see("end")
        except queue.Empty:
            pass
        self.after(100, self._drain_log_queue)

    def _browse(self):
        start_dir = Path.home() / "Downloads"
        if not start_dir.exists():
            start_dir = Path.home()
        path = filedialog.askopenfilename(
            parent=self,
            title="Selecciona Excel",
            initialdir=str(start_dir),
            filetypes=[("Excel", "*.xlsx"), ("Todos", "*.*")]
        )
        if path:
            self.file_var.set(path)

    def _start(self):
        if self.worker and self.worker.is_alive():
            return
        if not self.file_var.get():
            messagebox.showwarning("Falta archivo", "Selecciona un archivo .xlsx")
            return
        cfg = {
            "xlsx_path": self.file_var.get(),
            "window": self.window_var.get().strip() or "eleventa",
            "start_delay": float(self.start_delay_var.get()),
            "between_items": float(self.between_items_var.get()),
            "after_code_wait": float(self.after_code_wait_var.get()),
            "after_items_wait": float(self.after_items_wait_var.get()),
            "write_notes": bool(self.notes_var.get()),
            "finalize": bool(self.finalize_var.get()),
            "debug": bool(self.debug_var.get()),
            "common_prefix": self.common_prefix_var.get().strip() or "COMUN",
            "wholesale_all": bool(self.mayoreo_var.get()),
        }
        self._log("Preparando ejecución...")
        self.stop_flag.clear()
        self.pb.start(12)
        self.start_btn.configure(state="disabled")
        self.stop_btn.configure(state="normal")
        self.worker = threading.Thread(target=self._run_thread, args=(cfg,))
        self.worker.daemon = True
        self.worker.start()

    def _run_thread(self, cfg):
        try:
            run_sale(cfg, self._log, self.stop_flag)
        except Exception as e:
            self._log(f"[ERROR] {e}")
            messagebox.showerror("Error", str(e), parent=self)
        finally:
            self.pb.stop()
            self.start_btn.configure(state="normal")
            self.stop_btn.configure(state="disabled")

    def _stop(self):
        if self.worker and self.worker.is_alive():
            self._log("Cancelando...")
            self.stop_flag.set()

def main():
    app = App()
    app.mainloop()

if __name__ == "__main__":
    main()
